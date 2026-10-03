"""Build per-snapshot world border data + timeline text for the history map.
Inputs: aourednik/historical-basemaps geojson dir (GPL-3.0) and World_History xlsx.
Usage: python3 build_data.py <geojson_dir> <xlsx> <out_dir>"""
import json, sys, re, os
import openpyxl
from shapely.geometry import shape, Point
from shapely.ops import unary_union
from shapely.geometry.polygon import orient

GEO, XLSX, OUT = sys.argv[1:4]

# spreadsheet column (1-based) -> border file used for that snapshot
SNAP_FILES = {3:'bc4000',4:'bc3000',5:'bc1500',6:'bc700',7:'bc500',8:'bc400',9:'bc323',10:'bc200',11:'bc100',
  12:'bc1',13:'bc1',14:'100',15:'200',16:'300',17:'400',18:'500',19:'600',20:'700',21:'800',22:'900',
  23:'1000',24:'1100',25:'1279',26:'1300',27:'1492',28:'1530',29:'1650',30:'1715',31:'1783',32:'1815',
  33:'1815',34:'1878',35:'1880',36:'1900',37:'1914',38:'1920',39:'1938',40:'1945',41:'1960',42:'1960',
  43:'1960',44:'1960',45:'1994',46:'2000',47:'2010',48:'2010'}

# numeric year span of each column (BC negative)
SPANS = [(-10000,-3001),(-3000,-2001),(-2000,-1001),(-1000,-600),(-599,-500),(-499,-400),(-399,-300),(-299,-200),
  (-199,-100),(-99,-1),(1,99)] + [(c,c+99) for c in range(100,1700,100)] + [(1700,1749),(1750,1799),(1800,1824),
  (1825,1849),(1850,1874),(1875,1899)] + [(d,d+9) for d in range(1900,2030,10)]

REGIONS = {  # spreadsheet row -> anchor cities (lon, lat)
 'UK': [(-0.13,51.5),(-3.2,55.95),(-6.26,53.35)],
 'Scandinavia': [(18.07,59.33),(10.75,59.91),(12.57,55.68),(24.94,60.17)],
 'Spain & Portugal': [(-3.7,40.42),(-9.14,38.72),(-5.98,37.39)],
 'France': [(2.35,48.86),(4.84,45.76),(-1.68,48.11),(-0.58,44.84)],
 'Italy/Rome': [(12.5,41.9),(9.19,45.46),(13.36,38.12),(14.27,40.85),(9.11,39.22)],
 'Germany': [(13.4,52.52),(8.68,50.11),(11.58,48.14)],
 'Central Europe': [(21.01,52.23),(14.42,50.09),(16.37,48.21),(19.04,47.5)],
 'Russia': [(37.62,55.75),(30.52,50.45),(49.1,55.8),(82.9,55.0),(104.3,52.3),(131.9,43.1)],
 'Balkans': [(23.73,37.98),(20.46,44.79),(23.32,42.7),(28.98,41.01)],
 'Anatolia': [(32.86,39.93),(27.14,38.42),(35.5,38.7)],
 'Israel': [(35.21,31.77)],
 'Levant': [(36.29,33.51),(35.5,33.89),(37.16,36.2)],
 'Arabia': [(39.83,21.42),(46.72,24.69),(44.21,15.35),(58.4,23.6)],
 'Mesopotamia': [(44.36,33.31),(43.13,36.34),(47.8,30.5)],
 'Persia': [(51.39,35.69),(51.67,32.65),(52.89,29.94)],
 'Egypt': [(31.24,30.04),(29.92,31.2),(32.64,25.69)],
 'North Africa': [(10.18,36.81),(3.06,36.75),(-6.84,34.02),(13.19,32.89)],
 'West Africa': [(-0.19,5.6),(-3.0,16.77),(3.38,6.52),(-17.44,14.69),(7.5,12.0)],
 'East Africa': [(38.74,9.03),(36.82,-1.29),(45.32,2.05),(39.2,-6.16),(32.56,15.5)],
 'South Africa': [(18.42,-33.92),(28.05,-26.2),(30.93,-20.27),(25.9,-24.65)],
 'Central Asia / Steppe': [(66.96,39.65),(76.89,43.24),(106.9,47.92),(71.43,51.13),(58.38,37.95)],
 'India': [(77.21,28.61),(72.88,19.08),(88.36,22.57),(80.27,13.08),(74.35,31.55)],
 'China': [(116.4,39.9),(108.94,34.34),(121.47,31.23),(113.26,23.13),(104.07,30.67),(91.1,29.65)],
 'Southeast Asia': [(100.5,13.76),(105.85,21.03),(106.85,-6.21),(103.86,13.36),(120.98,14.6),(96.15,16.87)],
 'Japan': [(139.69,35.69),(135.77,35.01)],
 'Korea': [(126.98,37.57),(125.75,39.03)],
 'Oceania / Islanders': [(151.21,-33.87),(174.76,-36.85),(-157.86,21.31),(147.18,-9.44),(115.86,-31.95),(133.9,-23.7)],
 'United States': [(-77.04,38.9),(-74.0,40.71),(-87.63,41.88),(-118.24,34.05),(-96.8,32.78),(-79.38,43.65),(-104.99,39.74)],
 'Mexico': [(-99.13,19.43),(-100.31,25.67),(-103.35,20.67)],
 'Caribbean': [(-82.38,23.13),(-69.93,18.49),(-76.79,18.0),(-72.34,18.54)],
 'Central America': [(-90.51,14.63),(-87.19,14.07),(-79.52,8.98),(-89.62,20.97)],
 'Brazil': [(-43.17,-22.91),(-47.88,-15.79),(-38.5,-12.97),(-60.0,-3.1)],
 'Andes': [(-77.04,-12.05),(-71.97,-13.53),(-78.47,-0.18),(-68.15,-16.5),(-74.07,4.71)],
 'Southern Tip': [(-58.38,-34.6),(-70.67,-33.45),(-56.16,-34.9),(-57.58,-25.26)],
}
WORLD = ['Key Events','Global Events','World Population','Science & Technology','Religion & Ideas']

def fmt(p): return round(p, 2)

def clean(geom, tol=0.05):
    geom = geom.simplify(tol, preserve_topology=True).buffer(0)
    polys = [geom] if geom.geom_type == 'Polygon' else list(getattr(geom, 'geoms', []))
    polys = [orient(p, sign=-1.0) for p in polys if p.geom_type == 'Polygon' and p.area > 0.004]
    if not polys: return None
    coords = []
    for p in polys:
        rings = [[[fmt(x), fmt(y)] for x, y, *_ in p.exterior.coords]]
        rings += [[[fmt(x), fmt(y)] for x, y, *_ in r.coords] for r in p.interiors]
        coords.append(rings)
    return {'type': 'MultiPolygon', 'coordinates': coords}

# ---------- spreadsheet ----------
wb = openpyxl.load_workbook(XLSX); ws = wb['Sheet1']
def txt(v): return v.strip() if isinstance(v, str) else (str(v) if v is not None else '')
years = {c: txt(ws.cell(2, c).value) for c in range(3, 49)}
cent = {}; last = ''
for c in range(3, 49):
    v = txt(ws.cell(3, c).value); last = v or last; cent[c] = last
rows = {txt(ws.cell(r, 2).value).split('\n')[0]: r for r in range(4, 44)}

def cells(row):
    out = {}; tag = None; tagfill = None
    for c in range(3, 49):
        cell = ws.cell(row, c); v = txt(cell.value)
        fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
        m = re.match(r'^\[([^\]]+)\]\s*', v)
        if m: tag = m.group(1); tagfill = fill; v = v[m.end():]
        held = tag if fill and (m or fill == tagfill) else None
        out[c] = {'text': v, 'held': held}
    return out

timeline = {n: cells(rows[n]) for n in WORLD + ['European Events'] + list(REGIONS)}
snapshots = []
for i, c in enumerate(range(3, 49)):
    y = years[c]
    if y.startswith('Before'): label = 'Before 3000 BC'
    elif c <= 12: label = y + ' BC'
    elif c == 13: label = 'AD 0s'
    else: label = y
    f = SNAP_FILES[c]
    by = ('c. %s BC' % f[2:]) if f.startswith('bc') else ('c. AD %s' % f if int(f) < 1000 else 'c. %s' % f)
    if f == 'bc1': by = 'c. AD 1'
    snapshots.append({
        'label': label, 'century': cent[c], 'span': SPANS[i], 'file': f, 'borders': by,
        'world': {n: timeline[n][c]['text'] for n in WORLD + ['European Events'] if timeline[n][c]['text']},
        'regions': {n: timeline[n][c] for n in REGIONS if timeline[n][c]['text'] or timeline[n][c]['held']}})
# extra border years used only by war/religion step views (no slider stop of their own)
EXTRA_FILES = ['bc2000', 'bc1000', 'bc300', '1200', '1400', '1500', '1600', '1700', '1800', '1930']
ALL_FILES = sorted(set(SNAP_FILES.values()) | set(EXTRA_FILES), key=lambda f: -int(f[2:]) if f.startswith('bc') else int(f))
json.dump({'snapshots': snapshots, 'regions': list(REGIONS), 'borderFiles': ALL_FILES},
          open(os.path.join(OUT, 'timeline.json'), 'w'), ensure_ascii=False)

# ---------- borders ----------
anchors = [(r, Point(p)) for r, ps in REGIONS.items() for p in ps]
order = list(REGIONS)
for f in ALL_FILES:
    d = json.load(open(os.path.join(GEO, 'world_%s.geojson' % f)))
    groups = {}
    for ft in d['features']:
        pr = ft['properties']; n = (pr.get('NAME') or '').strip()
        if not n or not ft.get('geometry'): continue
        try: g = shape(ft['geometry']).buffer(0)
        except Exception: continue
        groups.setdefault(n, {'g': [], 'sub': pr.get('SUBJECTO') or '', 'part': pr.get('PARTOF') or ''})['g'].append(g)
    feats = []
    for n, info in groups.items():
        g = unary_union(info['g'])
        if g.is_empty or g.bounds[3] < -58: continue
        geo = clean(g)
        if not geo: continue
        parts = list(g.geoms) if hasattr(g, 'geoms') else [g]
        big = max(parts, key=lambda p: p.area)
        rp = big.representative_point()
        inside = big.buffer(0.05)
        home = next((r for r, p in anchors if inside.contains(p) and p.distance(rp) < 8), None) \
            or min(anchors, key=lambda a: a[1].distance(big))[0]
        regs = {r for r, p in anchors if g.buffer(0.05).contains(p)}
        regs = [home] + sorted(regs - {home}, key=order.index)
        props = {'name': n, 'regions': regs}
        for k in ('sub', 'part'):
            v = info[k].replace('�', 'ó') if info[k] else ''  # source typo: Arag�n
            if v and v != n: props[k] = v
        feats.append({'type': 'Feature', 'properties': props, 'geometry': geo})
    json.dump({'type': 'FeatureCollection', 'features': feats},
              open(os.path.join(OUT, 'borders_%s.json' % f), 'w'), separators=(',', ':'), ensure_ascii=False)
    print(f, len(feats), os.path.getsize(os.path.join(OUT, 'borders_%s.json' % f)) // 1024, 'KB')

# ---------- plain land base (from modern borders) ----------
d = json.load(open(os.path.join(GEO, 'world_2010.geojson')))
land = unary_union([shape(ft['geometry']).buffer(0) for ft in d['features'] if ft.get('geometry')])
land = land.simplify(0.15, preserve_topology=True).buffer(0)
polys = [p for p in (land.geoms if hasattr(land, 'geoms') else [land]) if p.area > 0.3]
json.dump({'type': 'Feature', 'properties': {}, 'geometry': clean(unary_union(polys))},
          open(os.path.join(OUT, 'land.json'), 'w'), separators=(',', ':'))
